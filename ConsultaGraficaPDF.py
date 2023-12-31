""" 
Curso Python empresa de 'Lenguaje de Programación Python'

Autor: José Antonio Calvo López

Fecha: Noviembre 2023

"""

from firebase_admin import credentials, firestore
from tkinter import ttk, PhotoImage
from fpdf import FPDF
import docx
import firebase_admin
import tkinter as tk
import sys, os


# Usar el archivo json descargado
cred = credentials.Certificate(r'alpine-charge-403816-508bb90d3f34.json')
firebase_admin.initialize_app(cred)

# Inicializa la instancia de Firestore
db = firestore.client()

#Funcion para obtener todos los datos de una colección
def get_all_data(collection_id):
    documents = db.collection(collection_id).stream() # se puede usar get() pero hace más uso de memoria aunque la obtención de datos es inmediata
    data = []
    for doc in documents:
        doc_data = doc.to_dict()
        # Opcional, si quieres incluir el ID del documento
        doc_data['id'] = doc.id  
        data.append(doc_data)
    return data

def mostrar_alumnos():
    borrar_datos()
    todos_alumnos = get_all_data('alumnos')
    for alumno in todos_alumnos:
        texto.insert(tk.END, f"{alumno}\n\n")

def mostrar_profesores():
    borrar_datos()
    datos_profesor = get_all_data('profesor')
    for profesor in datos_profesor:
        texto.insert(tk.END, f"{profesor}\n\n")

def generar_pdf():
    # Extrae las líneas de texto del widget
    datos = texto.get("1.0", tk.END).splitlines()
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)

    for dato in datos:
        pdf.cell(200, 10, txt=dato, ln=True, align='C')

    pdf.output("documentoPDF.pdf")

def generar_word():
    datos = texto.get("1.0", tk.END).splitlines()
    doc = docx.Document()

    for dato in datos:
        doc.add_paragraph(dato)

    doc.save("documentoDocx.docx")

def borrar_datos():
    texto.delete('1.0', tk.END)

# Crear la ventana principal
ventana = tk.Tk()
ventana.geometry("650x540")
ventana.title("Información A y P")

ventana.eval('tk::PlaceWindow . center')
ventana.resizable(False, False)

# Obtener la ruta de acceso a los recursos incluidos en el archivo
ruta_recursos = getattr(sys, '_MEIPASS', os.path.dirname(os.path.abspath(__file__)))

# Cargar las imágenes
icono = PhotoImage(file=os.path.join(ruta_recursos, "Studium.png"))

# Crear botones
boton_alumnos = ttk.Button(ventana, text="Mostrar Alumnos", command=mostrar_alumnos)
boton_alumnos.pack()

boton_profesores = ttk.Button(ventana, text="Mostrar Profesores", command=mostrar_profesores)
boton_profesores.pack()

# Crear área de texto
texto = tk.Text(ventana)
texto.pack()

boton_borrar = ttk.Button(ventana, text="Borrar", command=borrar_datos)
boton_borrar.pack()


separacion = ttk.Label(ventana, text="Utilidades")
separacion.pack()

boton_pdf = ttk.Button(ventana, text="Generar PDF", command=generar_pdf)
boton_pdf.pack()

boton_word = ttk.Button(ventana, text="Generar Word", command=generar_word)
boton_word.pack()

# Ejecutar la ventana
ventana.iconphoto(True, icono)
ventana.mainloop()